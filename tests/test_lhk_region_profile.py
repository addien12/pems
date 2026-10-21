from datetime import date

import openpyxl

from app.import_service import ACTIVE, validate_workbook


def make_lhk_workbook(path):
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Produksi Tahun Ini"
    worksheet["B3"] = "TANGGAL: 2026-10-20"
    for day in range(1, 32):
        worksheet.cell(6, 4 + day, day)

    row = 7
    for estate, count in ACTIVE.items():
        for division_number in range(1, count + 1):
            worksheet.cell(row, 2, f"{estate}{division_number}")
            for day in range(1, 21):
                worksheet.cell(row, 4 + day, 10)
            row += 1

    workbook.save(path)
    workbook.close()


def test_lhk_region_profile_is_daily_year_to_date(tmp_path):
    path = tmp_path / "DB_LHK_PANEN_REGION_2026-10-20.xlsx"
    make_lhk_workbook(path)

    preview = validate_workbook(path, "actual")

    assert preview["source_system"] == "DB LHK REGION"
    assert preview["update_frequency"] == "daily"
    assert preview["report_scope"] == "year_to_date"
    assert preview["report_date"] == date(2026, 10, 20).isoformat()
    assert preview["division_count"] == 24
    assert preview["record_count"] == 24 * 20
    assert preview["total_ton"] == "4800.0000"
