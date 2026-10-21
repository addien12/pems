from pathlib import Path
from datetime import date, datetime
import openpyxl

path = Path("imports/original/DB_LHK_PANEN_REGION.xlsx")
if not path.is_file():
    raise SystemExit(f"File tidak ditemukan: {path.resolve()}")

active = {
    "MAGE": 6,
    "BMKE": 4,
    "BKCE": 4,
    "BBGE": 3,
    "BHPE": 4,
    "SPGE": 3,
}

wb = openpyxl.load_workbook(
    path, read_only=False, data_only=False
)

try:
    sheet_name = "Produksi Tahun Ini"
    if sheet_name not in wb.sheetnames:
        raise SystemExit(f"Sheet {sheet_name!r} tidak ditemukan")

    ws = wb[sheet_name]
    raw_date = ws["B3"].value

    if isinstance(raw_date, datetime):
        report_date = raw_date.date()
    elif isinstance(raw_date, date):
        report_date = raw_date
    elif isinstance(raw_date, str):
        text = raw_date.strip().split(":", 1)[-1].strip()
        try:
            report_date = date.fromisoformat(text[:10])
        except ValueError:
            report_date = None
    else:
        report_date = None

    print("File:", path.name)
    print("Sheet:", sheet_name)
    print("Dimensi:", ws.max_row, "x", ws.max_column)
    print("B3:", repr(raw_date))
    print("Tanggal laporan:", report_date)
    print("Header E6:AI6:", [
        ws.cell(6, c).value for c in range(5, 36)
    ])

    expected = {
        f"{estate}{n}"
        for estate, count in active.items()
        for n in range(1, count + 1)
    }
    rows = {}

    for row in range(1, ws.max_row + 1):
        value = ws.cell(row, 2).value
        if isinstance(value, str) and value.strip().upper() in expected:
            rows.setdefault(value.strip().upper(), []).append(row)

    print("\nPEMETAAN DIVISI")
    errors = []
    total = 0.0
    record_count = 0

    if not report_date:
        print("Tanggal laporan tidak valid; pemeriksaan harian dihentikan.")
    else:
        for code in sorted(expected):
            matches = rows.get(code, [])
            if len(matches) != 1:
                print(f"{code}: BARIS TIDAK VALID {matches}")
                continue

            row = matches[0]
            values = [
                ws.cell(row, 4 + day).value
                for day in range(1, report_date.day + 1)
            ]
            invalid = [
                (day, value)
                for day, value in enumerate(values, 1)
                if isinstance(value, bool)
                or not isinstance(value, (int, float))
                or value < 0
            ]

            numeric_total = sum(
                value for value in values
                if isinstance(value, (int, float))
                and not isinstance(value, bool)
                and value >= 0
            )
            print(
                f"{code}: baris={row}, "
                f"angka={len(values)-len(invalid)}/{len(values)}, "
                f"total={numeric_total:.4f}, "
                f"hari1-3={values[:3]}"
            )

            if invalid:
                errors.append((code, invalid[:3]))
            else:
                record_count += len(values)
                total += numeric_total

    print("\nKode hilang:", sorted(expected - set(rows)))
    print("Kode ganda:", {
        code: positions for code, positions in rows.items()
        if len(positions) != 1
    })
    print("Jumlah record valid:", record_count)
    print("Total sementara:", round(total, 4), "ton")
    print("Kesalahan nilai:", errors[:10])

finally:
    wb.close()
