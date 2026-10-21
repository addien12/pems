from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
import argparse
import hashlib
import os
import shutil
import sys

import openpyxl
import psycopg

SHEET = "Produksi Tahun Ini"
TON_QUANTUM = Decimal("0.0001")
ACTIVE = {
    "MAGE": 6,
    "BMKE": 4,
    "BKCE": 4,
    "BBGE": 3,
    "BHPE": 4,
    "SPGE": 3,
}


def parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        value = value.strip().split(":", 1)[-1].strip()
        try:
            return date.fromisoformat(value[:10])
        except ValueError:
            return None
    return None


def sha256_file(path):
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def read_source(path):
    workbook = openpyxl.load_workbook(
        path, read_only=True, data_only=False
    )
    try:
        if SHEET not in workbook.sheetnames:
            raise ValueError(f"Sheet {SHEET!r} tidak ditemukan")

        ws = workbook[SHEET]
        report_date = parse_date(ws["B3"].value)

        if not report_date:
            raise ValueError(
                f"Tanggal laporan B3 tidak valid: {ws['B3'].value!r}"
            )

        headers = [ws.cell(6, c).value for c in range(5, 36)]
        if headers != list(range(1, 32)):
            raise ValueError("Header hari E6:AI6 bukan 1 sampai 31")

        expected = {
            f"{estate}{n}"
            for estate, count in ACTIVE.items()
            for n in range(1, count + 1)
        }
        rows = {}

        for r in range(1, ws.max_row + 1):
            value = ws.cell(r, 2).value
            if isinstance(value, str):
                code = value.strip().upper()
                if code in expected:
                    rows.setdefault(code, []).append(r)

        missing = sorted(expected - set(rows))
        duplicate = {
            code: positions
            for code, positions in rows.items()
            if len(positions) != 1
        }
        if missing or duplicate:
            raise ValueError(
                f"Kode divisi tidak valid: hilang={missing}; "
                f"ganda={duplicate}"
            )

        records = []
        errors = []

        for code in sorted(expected):
            row = rows[code][0]

            for day in range(1, report_date.day + 1):
                value = ws.cell(row, 4 + day).value

                if (
                    isinstance(value, bool)
                    or not isinstance(value, (int, float))
                ):
                    errors.append(
                        f"{code}, hari {day}: bukan angka ({value!r})"
                    )
                    continue

                amount = Decimal(str(value)).quantize(TON_QUANTUM)
                if not amount.is_finite() or amount < 0:
                    errors.append(
                        f"{code}, hari {day}: nilai tidak valid ({value})"
                    )
                    continue

                records.append((
                    code,
                    row,
                    date(report_date.year, report_date.month, day),
                    amount,
                ))

        if errors:
            raise ValueError(
                "Validasi gagal:\n- " + "\n- ".join(errors[:30])
            )

        if len(records) != len(expected) * report_date.day:
            raise ValueError(
                f"Jumlah record tidak sesuai: {len(records)}"
            )

        return report_date, records
    finally:
        workbook.close()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("file", help="Path workbook DB LHK")
    parser.add_argument(
        "--commit",
        action="store_true",
        help="Simpan hasil impor ke PostgreSQL",
    )
    args = parser.parse_args()

    path = Path(args.file).expanduser().resolve()
    if not path.is_file():
        raise ValueError(f"File tidak ditemukan: {path}")

    database_url = os.getenv("DATABASE_URL")
    if not database_url:
        raise ValueError("DATABASE_URL belum diatur")

    file_hash = sha256_file(path)
    report_date, records = read_source(path)
    total = sum((r[3] for r in records), Decimal("0"))

    print("===== VALIDASI REALISASI PRODUKSI =====")
    print(f"File            : {path.name}")
    print(f"Sheet           : {SHEET}")
    print(f"SHA-256         : {file_hash}")
    print(f"Tanggal laporan : {report_date}")
    print("Jumlah divisi   : 24")
    print(f"Jumlah record   : {len(records)}")
    print(f"Total aktual    : {total:,.4f} ton")
    print(f"Mode            : {'COMMIT' if args.commit else 'UJI COBA'}")

    with psycopg.connect(database_url) as conn:
        with conn.cursor() as cur:
            expected_estates = set(ACTIVE)

            cur.execute(
                "SELECT code FROM estates WHERE is_active = TRUE"
            )
            actual_estates = {r[0] for r in cur.fetchall()}
            if actual_estates != expected_estates:
                raise ValueError(
                    f"Estate aktif tidak sesuai: {sorted(actual_estates)}"
                )

            expected_divisions = {
                f"{estate}{n}"
                for estate, count in ACTIVE.items()
                for n in range(1, count + 1)
            }
            cur.execute("""
                SELECT d.code
                FROM divisions d
                JOIN estates e ON e.id = d.estate_id
                WHERE d.is_active = TRUE AND e.is_active = TRUE
            """)
            actual_divisions = {r[0] for r in cur.fetchall()}

            if actual_divisions != expected_divisions:
                raise ValueError("Daftar divisi aktif tidak sesuai")

            cur.execute("""
                SELECT id FROM import_batches
                WHERE source_sha256 = %s AND source_sheet = %s
                LIMIT 1
            """, (file_hash, SHEET))
            existing = cur.fetchone()

            if existing:
                print(
                    f"Batch identik sudah ada: {existing[0]}. "
                    "Tidak ada perubahan."
                )
                return

            if not args.commit:
                print("UJI COBA BERHASIL. Database tidak diubah.")
                print("Jika hasil sudah benar, jalankan dengan --commit.")
                return

            archive_dir = Path.home() / "pems" / "imports" / "original"
            archive_dir.mkdir(parents=True, exist_ok=True)
            archive_path = (
                archive_dir /
                f"{path.stem}_{file_hash[:12]}{path.suffix.lower()}"
            )

            if path != archive_path and not archive_path.exists():
                shutil.copy2(path, archive_path)

            cur.execute("""
                INSERT INTO import_batches
                    (source_filename, source_sha256, source_sheet,
                     report_date, imported_by, notes)
                VALUES (%s, %s, %s, %s, %s, %s)
                RETURNING id
            """, (
                path.name,
                file_hash,
                SHEET,
                report_date,
                os.getenv("PEMS_IMPORT_USER", "termux-cli"),
                "Produksi Tahun Ini; realisasi harian per divisi",
            ))
            batch_id = cur.fetchone()[0]

            cur.execute("""
                SELECT d.code, d.id
                FROM divisions d
                JOIN estates e ON e.id = d.estate_id
                WHERE d.is_active = TRUE AND e.is_active = TRUE
            """)
            division_ids = dict(cur.fetchall())

            cur.executemany("""
                INSERT INTO production_actuals
                    (import_batch_id, division_id, actual_date,
                     actual_ton, source_row)
                VALUES (%s, %s, %s, %s, %s)
            """, [
                (batch_id, division_ids[code], day, amount, row)
                for code, row, day, amount in records
            ])

            cur.execute("""
                INSERT INTO audit_logs
                    (actor, action, entity_type, entity_id, details)
                VALUES (%s, %s, %s, %s, %s::jsonb)
            """, (
                os.getenv("PEMS_IMPORT_USER", "termux-cli"),
                "IMPORT_PRODUKSI_AKTUAL",
                "import_batch",
                str(batch_id),
                (
                    '{"status":"success","active_divisions":24,'
                    '"record_count":%d}' % len(records)
                ),
            ))

    print("IMPORT AKTUAL BERHASIL")
    print(f"Batch ID        : {batch_id}")
    print(f"Record tersimpan: {len(records)}")
    print(f"Total aktual    : {total:,.4f} ton")


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        sys.exit(f"IMPORT DIBATALKAN: {error}")
