from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
import argparse
import hashlib
import os
import shutil
import sys

import openpyxl
import psycopg

SHEET = "Budget Produksi"
ACTIVE = {
    "MAGE": 6,
    "BMKE": 4,
    "BKCE": 4,
    "BBGE": 3,
    "BHPE": 4,
    "SPGE": 3,
}


def parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        value = value.strip().split(":", 1)[-1].strip()
        try:
            return date.fromisoformat(value[:10])
        except ValueError:
            return None
    return None


def digest_file(path):
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def read_source(path):
    workbook = openpyxl.load_workbook(
        path, read_only=True, data_only=False
    )
    try:
        if SHEET not in workbook.sheetnames:
            raise ValueError(f"Sheet {SHEET!r} tidak ditemukan")

        sheet = workbook[SHEET]
        report_date = parse_date(sheet["B3"].value)
        if not report_date:
            raise ValueError(
                f"Tanggal laporan B3 tidak terbaca: "
                f"{sheet['B3'].value!r}"
            )

        if [sheet.cell(6, col).value for col in range(5, 36)] != list(
            range(1, 32)
        ):
            raise ValueError("Header hari E6:AI6 bukan 1 sampai 31")

        expected = {
            f"{estate}{number}"
            for estate, count in ACTIVE.items()
            for number in range(1, count + 1)
        }
        rows = {}

        for row_number in range(1, sheet.max_row + 1):
            value = sheet.cell(row_number, 2).value
            if isinstance(value, str):
                code = value.strip().upper()
                if code in expected:
                    rows.setdefault(code, []).append(row_number)

        missing = sorted(expected - set(rows))
        duplicates = {
            code: positions
            for code, positions in rows.items()
            if len(positions) != 1
        }
        if missing or duplicates:
            raise ValueError(
                f"Kode divisi tidak valid. Hilang={missing}; "
                f"ganda={duplicates}"
            )

        records = []
        errors = []

        # Hanya membaca tanggal 1 sampai tanggal laporan.
        for code in sorted(expected):
            row_number = rows[code][0]
            for day in range(1, report_date.day + 1):
                value = sheet.cell(row_number, 4 + day).value

                if (
                    isinstance(value, bool)
                    or not isinstance(value, (int, float))
                ):
                    errors.append(
                        f"{code}, baris {row_number}, hari {day}: "
                        f"bukan angka ({value!r})"
                    )
                    continue

                amount = Decimal(str(value))
                if not amount.is_finite() or amount < 0:
                    errors.append(
                        f"{code}, hari {day}: nilai tidak valid ({value})"
                    )
                    continue

                records.append(
                    (
                        code,
                        row_number,
                        date(report_date.year, report_date.month, day),
                        amount,
                    )
                )

        if errors:
            raise ValueError(
                "Validasi anggaran gagal:\n- " + "\n- ".join(errors[:30])
            )

        expected_count = len(expected) * report_date.day
        if len(records) != expected_count:
            raise ValueError(
                f"Record {len(records)}, seharusnya {expected_count}"
            )

        return report_date, records
    finally:
        workbook.close()


def verify_reference_data(cursor):
    cursor.execute(
        "SELECT code FROM estates WHERE is_active = TRUE"
    )
    estates = {row[0] for row in cursor.fetchall()}

    if estates != set(ACTIVE):
        raise ValueError(
            f"Referensi estate aktif tidak sesuai. "
            f"Database={sorted(estates)}, "
            f"seharusnya={sorted(ACTIVE)}"
        )

    cursor.execute(
        "SELECT d.code FROM divisions d "
        "JOIN estates e ON e.id = d.estate_id "
        "WHERE d.is_active = TRUE AND e.is_active = TRUE"
    )
    divisions = {row[0] for row in cursor.fetchall()}

    expected = {
        f"{estate}{number}"
        for estate, count in ACTIVE.items()
        for number in range(1, count + 1)
    }
    if divisions != expected:
        raise ValueError(
            f"Referensi divisi aktif tidak sesuai. "
            f"Hilang={sorted(expected - divisions)}, "
            f"tambahan={sorted(divisions - expected)}"
        )

    cursor.execute(
        "SELECT code, id FROM divisions "
        "WHERE is_active = TRUE"
    )
    return dict(cursor.fetchall())


def main():
    parser = argparse.ArgumentParser(
        description="Impor anggaran harian DB LHK ke PEMS"
    )
    parser.add_argument("file", help="Path file DB LHK .xlsx")
    parser.add_argument(
        "--commit",
        action="store_true",
        help="Simpan impor ke database; tanpa ini hanya uji coba",
    )
    args = parser.parse_args()

    path = Path(args.file).expanduser().resolve()
    if not path.is_file():
        raise ValueError(f"File tidak ditemukan: {path}")

    database_url = os.getenv("DATABASE_URL")
    if not database_url:
        raise ValueError(
            "DATABASE_URL belum diatur. "
            "Contoh lokal: export DATABASE_URL='postgresql:///pems_db'"
        )

    file_hash = digest_file(path)
    report_date, records = read_source(path)
    total = sum((record[3] for record in records), Decimal("0"))

    print("===== HASIL VALIDASI SUMBER =====")
    print(f"File             : {path.name}")
    print(f"SHA-256          : {file_hash}")
    print(f"Tanggal laporan  : {report_date}")
    print(f"Sheet            : {SHEET}")
    print(f"Jumlah divisi    : 24")
    print(f"Jumlah hari      : {report_date.day}")
    print(f"Jumlah record    : {len(records)}")
    print(f"Total budget     : {total:,.4f} ton")
    print(f"Mode             : {'SIMPAN' if args.commit else 'UJI COBA'}")

    archive_dir = Path.home() / "pems" / "imports" / "original"
    archive_path = (
        archive_dir / f"{path.stem}_{file_hash[:12]}{path.suffix.lower()}"
    )

    with psycopg.connect(database_url) as connection:
        with connection.cursor() as cursor:
            division_ids = verify_reference_data(cursor)

            cursor.execute(
                "SELECT id FROM import_batches "
                "WHERE source_sha256 = %s AND source_sheet = %s "
                "LIMIT 1",
                (file_hash, SHEET),
            )
            existing = cursor.fetchone()

            if existing:
                print(
                    f"\nFile identik sudah terdaftar pada batch "
                    f"{existing[0]}. Tidak ada perubahan."
                )
                return

            if not args.commit:
                print("\nUJI COBA SELESAI. Database tidak diubah.")
                print("Jika hasilnya benar, jalankan ulang dengan --commit.")
                return

            archive_dir.mkdir(parents=True, exist_ok=True)
            if path != archive_path and not archive_path.exists():
                shutil.copy2(path, archive_path)

            cursor.execute(
                """
                INSERT INTO import_batches
                    (source_filename, source_sha256, source_sheet,
                     report_date, imported_by, notes)
                VALUES (%s, %s, %s, %s, %s, %s)
                RETURNING id
                """,
                (
                    path.name,
                    file_hash,
                    SHEET,
                    report_date,
                    os.getenv("PEMS_IMPORT_USER", "termux-cli"),
                    "Budget Produksi; 24 divisi; hari 1 sampai tanggal laporan",
                ),
            )
            batch_id = cursor.fetchone()[0]

            cursor.executemany(
                """
                INSERT INTO production_budgets
                    (import_batch_id, division_id, budget_date,
                     budget_ton, source_row)
                VALUES (%s, %s, %s, %s, %s)
                """,
                [
                    (
                        batch_id,
                        division_ids[code],
                        budget_date,
                        amount,
                        source_row,
                    )
                    for code, source_row, budget_date, amount in records
                ],
            )

            cursor.execute(
                """
                INSERT INTO audit_logs
                    (actor, action, entity_type, entity_id, details)
                VALUES (%s, %s, %s, %s, %s::jsonb)
                """,
                (
                    os.getenv("PEMS_IMPORT_USER", "termux-cli"),
                    "IMPORT_BUDGET_PRODUKSI",
                    "import_batch",
                    str(batch_id),
                    (
                        '{"status":"success","active_divisions":24,'
                        '"record_count":%d}' % len(records)
                    ),
                ),
            )

        # Commit dilakukan oleh context psycopg setelah blok berhasil.
    print("\nIMPORT BERHASIL")
    print(f"Batch ID         : {batch_id}")
    print(f"Arsip            : {archive_path}")
    print(f"Record tersimpan : {len(records)}")
    print(f"Total budget     : {total:,.4f} ton")


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        sys.exit(f"\nIMPORT DIBATALKAN: {error}")
