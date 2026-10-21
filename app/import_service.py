from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
import hashlib
import os

import openpyxl
import psycopg


ACTIVE = {
    "MAGE": 6,
    "BMKE": 4,
    "BKCE": 4,
    "BBGE": 3,
    "BHPE": 4,
    "SPGE": 3,
}

MAX_WORKSHEET_ROWS = 1000
MAX_WORKSHEET_COLUMNS = 40
TON_QUANTUM = Decimal("0.0001")

SHEETS = {
    "budget": ("Budget Produksi", "production_budgets", "budget_date", "budget_ton",
               "IMPORT_BUDGET_PRODUKSI"),
    "actual": ("Produksi Tahun Ini", "production_actuals", "actual_date", "actual_ton",
               "IMPORT_PRODUKSI_AKTUAL"),
}
IMPORT_PROFILES = {
    "budget": {
        "source_system": "PEMS Budget Workbook",
        "frequency": "periodic",
        "scope": "period",
    },
    "actual": {
        "source_system": "DB LHK REGION",
        "frequency": "daily",
        "scope": "year_to_date",
    },
}


def parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            return date.fromisoformat(value.strip().split(":")[-1].strip()[:10])
        except ValueError:
            return None
    return None


def validate_workbook(path: Path, kind: str):
    if kind not in SHEETS:
        raise ValueError("Jenis impor harus budget atau actual.")

    sheet_name = SHEETS[kind][0]
    profile = IMPORT_PROFILES[kind]
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)

    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Sheet {sheet_name!r} tidak ditemukan.")

        ws = workbook[sheet_name]
        if ws.max_row > MAX_WORKSHEET_ROWS or ws.max_column > MAX_WORKSHEET_COLUMNS:
            raise ValueError(
                "Ukuran workbook melampaui batas yang diizinkan "
                f"({MAX_WORKSHEET_ROWS} baris, {MAX_WORKSHEET_COLUMNS} kolom)."
            )
        report_date = parse_date(ws["B3"].value)
        if not report_date:
            raise ValueError(f"Tanggal laporan di B3 tidak valid: {ws['B3'].value!r}")

        headers = [ws.cell(6, col).value for col in range(5, 36)]
        if headers != list(range(1, 32)):
            raise ValueError("Header hari E6:AI6 harus berisi angka 1 sampai 31.")

        expected = {
            f"{estate}{n}"
            for estate, count in ACTIVE.items()
            for n in range(1, count + 1)
        }

        locations = {}
        for row in range(1, ws.max_row + 1):
            value = ws.cell(row, 2).value
            if isinstance(value, str):
                code = value.strip().upper()
                if code in expected:
                    locations.setdefault(code, []).append(row)

        missing = sorted(expected - set(locations))
        duplicates = {k: v for k, v in locations.items() if len(v) != 1}
        if missing or duplicates:
            raise ValueError(
                f"Kode divisi tidak sesuai. Hilang={missing}; duplikat={duplicates}"
            )

        records = []
        errors = []
        for code in sorted(expected):
            row = locations[code][0]
            for day in range(1, report_date.day + 1):
                value = ws.cell(row, 4 + day).value
                if isinstance(value, bool) or not isinstance(value, (int, float)):
                    errors.append(f"{code}, baris {row}, hari {day}: bukan angka ({value!r})")
                    continue

                amount = Decimal(str(value)).quantize(TON_QUANTUM)
                if not amount.is_finite() or amount < 0:
                    errors.append(f"{code}, hari {day}: nilai tidak valid ({value!r})")
                    continue

                records.append((
                    code, row,
                    date(report_date.year, report_date.month, day),
                    amount,
                ))

        if errors:
            raise ValueError("Validasi gagal:\n- " + "\n- ".join(errors[:30]))

        expected_count = len(expected) * report_date.day
        if len(records) != expected_count:
            raise ValueError(f"Jumlah record {len(records)}; seharusnya {expected_count}.")

        return {
            "kind": kind,
            "sheet": sheet_name,
            "source_system": profile["source_system"],
            "update_frequency": profile["frequency"],
            "report_scope": profile["scope"],
            "report_date": report_date.isoformat(),
            "division_count": len(expected),
            "record_count": len(records),
            "total_ton": str(sum((r[3] for r in records), Decimal("0"))),
            "records": records,
        }
    finally:
        workbook.close()


def sha256_file(path: Path):
    digest = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def save_import(
    path: Path,
    kind: str,
    actor: str = "web-user",
    source_filename: str | None = None,
):
    filename = Path(source_filename or path.name).name
    if kind not in SHEETS:
        raise ValueError("Jenis impor tidak dikenal.")

    preview = validate_workbook(path, kind)
    file_hash = sha256_file(path)
    sheet, table, date_col, amount_col, action = SHEETS[kind]
    database_url = os.getenv("DATABASE_URL", "postgresql:///pems_db")

    with psycopg.connect(database_url) as conn:
        with conn.cursor() as cur:
            cur.execute(
                "SELECT id FROM import_batches WHERE source_sha256=%s AND source_sheet=%s",
                (file_hash, sheet),
            )
            existing = cur.fetchone()
            if existing:
                raise ValueError(f"File identik sudah diimpor pada batch {existing[0]}.")

            cur.execute(
                """
                SELECT d.code, d.id
                FROM divisions d
                JOIN estates e ON e.id = d.estate_id
                WHERE d.is_active = TRUE AND e.is_active = TRUE
                """
            )
            division_ids = dict(cur.fetchall())

            expected = {
                f"{estate}{n}"
                for estate, count in ACTIVE.items()
                for n in range(1, count + 1)
            }
            if set(division_ids) != expected:
                raise ValueError("Daftar divisi aktif database tidak sesuai.")

            cur.execute("""
                INSERT INTO import_batches
                    (source_filename, source_sha256, source_sheet,
                     report_date, imported_by, notes)
                VALUES (%s, %s, %s, %s, %s, %s)
                RETURNING id
            """, (
                filename, file_hash, sheet, preview["report_date"], actor,
                (
                    f"Sumber {preview['source_system']}; "
                    f"scope {preview['report_scope']}; "
                    f"frekuensi {preview['update_frequency']}; "
                    f"{preview['record_count']} record"
                ),
            ))
            batch_id = cur.fetchone()[0]

            insert_sql = f"""
                INSERT INTO {table}
                    (import_batch_id, division_id, {date_col}, {amount_col}, source_row)
                VALUES (%s, %s, %s, %s, %s)
            """
            cur.executemany(insert_sql, [
                (batch_id, division_ids[code], day, amount, source_row)
                for code, source_row, day, amount in preview["records"]
            ])

            import json
            cur.execute("""
                INSERT INTO audit_logs
                    (actor, action, entity_type, entity_id, details)
                VALUES (%s, %s, %s, %s, %s::jsonb)
            """, (
                actor, action, "import_batch", str(batch_id),
                json.dumps({
                    "status": "success",
                    "filename": filename,
                    "sha256": file_hash,
                    "record_count": preview["record_count"],
                    "report_date": preview["report_date"],
                    "source_system": preview["source_system"],
                    "report_scope": preview["report_scope"],
                }),
            ))

    return {
        "batch_id": batch_id,
        "filename": filename,
        "sha256": file_hash,
        "sheet": sheet,
        "source_system": preview["source_system"],
        "update_frequency": preview["update_frequency"],
        "report_scope": preview["report_scope"],
        "report_date": preview["report_date"],
        "record_count": preview["record_count"],
        "total_ton": preview["total_ton"],
    }
